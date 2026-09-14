"""Collect every tuning trial into one readable spreadsheet.

Inputs:  research/results/tuning/<model>/<dataset>/<recipe>/seed<n>/
             baseline.json, <knob>/<value>.json, winner.json
         (written by research.runners.tuning).
Outputs: research/results/tuning_summary.xlsx  (sheets: Trials, Winners, Coverage)
         research/results/tuning_summary.csv   (the Trials sheet, flat)

Trials the grid in tuning_config.TUNE_ORDER calls for but that are not on disk
are listed with status "missing", so an unfinished cell shows its gaps instead of
looking complete.

    python -m research.analysis.aggregate_tuning
"""

import json

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..lib.tuning_config import TUNE_ORDER
from ..paths import RESULTS_DIR

IN_DIR = RESULTS_DIR / "tuning"
OUT_XLSX = RESULTS_DIR / "tuning_summary.xlsx"
OUT_CSV = RESULTS_DIR / "tuning_summary.csv"

MODELS = ["bpr", "lightgcn", "neumf"]
DATASETS = ["musical", "baby", "cellphone", "healthcare"]
ARMS = ["uniform", "causal"]
SEEDS = [42, 123, 2026]
METRICS = [("HR@20", "HitRatio@20"), ("NDCG@20", "NDCG@20"), ("Recall@20", "Recall@20")]

WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
MISSING_FILL = PatternFill("solid", fgColor="E7E6E6")
STALE_FILL = PatternFill("solid", fgColor="FFEB9C")


def _value_slug(value) -> str:
    # Same encoding as runners/tuning.py:_value_slug -- copied rather than
    # imported, because importing the runner pulls in cornac and torch.
    s = repr(value) if isinstance(value, bool) else str(value)
    return (s.replace(".", "p").replace("-", "m").replace("+", "")
             .replace("(", "").replace(")", "").replace(",", "_").replace(" ", ""))


def _load(path):
    with open(path, encoding="utf-8-sig") as f:
        return json.load(f)


def _base_model(model_dir: str) -> str:
    """`neumf-pretrain` -> `neumf`, the key TUNE_ORDER is written under."""
    return model_dir.split("-")[0]


def _same_config(a: dict, b: dict) -> bool:
    if a.keys() != b.keys():
        return False
    for k in a:
        try:
            if abs(float(a[k]) - float(b[k])) > 1e-12 * max(1.0, abs(float(a[k]))):
                return False
        except (TypeError, ValueError):
            if a[k] != b[k]:
                return False
    return True


def _cells():
    for cell in sorted(IN_DIR.glob("*/*/*/seed*")):
        if cell.is_dir():
            model_dir, dataset, arm = cell.parts[-4:-1]
            yield model_dir, dataset, arm, int(cell.name[len("seed"):]), cell


def _trial_row(ident, knob, value, payload, winner_cfg):
    row = dict(ident, knob=knob, value=value)
    if payload is None:
        row["status"] = "missing"
        return row
    cfg = payload.get("config", {})
    val = payload.get("val_metrics")
    test = payload.get("test_metrics", {})
    # Pre-2026-09-11 files have no val_metrics: they were selected on test and
    # the tuner treats them as stale.
    row["status"] = "done" if val else "stale"
    row["winner"] = "yes" if winner_cfg is not None and _same_config(cfg, winner_cfg) else ""
    for disp, key in METRICS:
        row[f"val {disp}"] = (val or {}).get(key)
    for disp, key in METRICS:
        row[f"test {disp}"] = test.get(key)
    row["train (s)"] = test.get("Train (s)")
    row["wall (h)"] = (payload["wall_seconds"] / 3600) if "wall_seconds" in payload else None
    for k, v in cfg.items():
        row[f"cfg.{k}"] = v
    return row


def collect():
    trials, winners = [], []
    for model_dir, dataset, arm, seed, cell in _cells():
        ident = dict(model=model_dir, dataset=dataset, arm=arm, seed=seed)
        winner_path = cell / "winner.json"
        winner = _load(winner_path) if winner_path.exists() else None
        winner_cfg = winner["config"] if winner else None
        start = len(trials)

        baseline = cell / "baseline.json"
        trials.append(_trial_row(ident, "baseline", "",
                                 _load(baseline) if baseline.exists() else None, winner_cfg))
        seen = set()
        for knob, grid in TUNE_ORDER.get(_base_model(model_dir), []):
            for value in grid:
                path = cell / knob / f"{_value_slug(value)}.json"
                seen.add(path)
                trials.append(_trial_row(ident, knob, value,
                                         _load(path) if path.exists() else None, winner_cfg))
        # Trials for knobs no longer in TUNE_ORDER (e.g. lambda_reg before the trim).
        for path in sorted(cell.glob("*/*.json")):
            if path not in seen:
                row = _trial_row(ident, path.parent.name, path.stem, _load(path), winner_cfg)
                row["status"] = "off-grid"
                trials.append(row)

        n_expected = 1 + sum(len(g) for _, g in TUNE_ORDER.get(_base_model(model_dir), []))
        n_done = sum(1 for r in trials[start:] if r["status"] == "done")
        w = dict(ident, trials=f"{n_done}/{n_expected}", finished="yes" if winner else "")
        if winner:
            w["val NDCG@20"] = winner.get("metric")
            w["test NDCG@20"] = winner.get("test_metric")
            w["select_split"] = winner.get("select_split")
            probes = (winner.get("users_run") or {}).get("probes") or {}
            w["counterfactual_rate"] = probes.get("counterfactual_rate")
            w["collision_rate"] = probes.get("collision_rate")
            for k, v in winner_cfg.items():
                w[f"cfg.{k}"] = v
        winners.append(w)
    return pd.DataFrame(trials), pd.DataFrame(winners)


def add_arm_gap(winners: pd.DataFrame) -> pd.DataFrame:
    """causal vs uniform winner, as % of uniform, where both arms finished."""
    key = ["model", "dataset", "seed"]
    uni = (winners[winners["arm"] == "uniform"]
           .set_index(key)[["val NDCG@20", "test NDCG@20"]])
    gaps = []
    for _, r in winners.iterrows():
        k = (r["model"], r["dataset"], r["seed"])
        if r["arm"] != "causal" or k not in uni.index or pd.isna(r.get("val NDCG@20")):
            gaps.append((None, None))
            continue
        u = uni.loc[k]
        gaps.append(tuple(
            (r[m] / u[m] - 1) * 100 if pd.notna(u[m]) and u[m] else None
            for m in ["val NDCG@20", "test NDCG@20"]))
    winners["causal vs uniform val %"] = [g[0] for g in gaps]
    winners["causal vs uniform test %"] = [g[1] for g in gaps]
    return winners


def coverage(winners: pd.DataFrame) -> pd.DataFrame:
    """One row per model/dataset/arm, one column per seed: trials done / expected."""
    models = sorted(set(winners["model"]) | set(MODELS))
    datasets = [d for d in DATASETS if d in set(winners["dataset"])] or DATASETS
    lookup = {(r["model"], r["dataset"], r["arm"], r["seed"]): r
              for _, r in winners.iterrows()}
    rows = []
    for m in models:
        for d in datasets:
            for a in ARMS:
                row = dict(model=m, dataset=d, arm=a)
                for s in SEEDS:
                    r = lookup.get((m, d, a, s))
                    row[f"seed {s}"] = ("" if r is None
                                        else f"{r['trials']} ✓" if r["finished"] else r["trials"])
                rows.append(row)
    return pd.DataFrame(rows)


def _order(trials: pd.DataFrame) -> pd.DataFrame:
    knob_rank = {"baseline": 0}
    for knobs in TUNE_ORDER.values():
        for i, (k, _) in enumerate(knobs, start=1):
            knob_rank.setdefault(k, i)
    trials = trials.assign(
        _m=trials["model"],
        _d=trials["dataset"].map({d: i for i, d in enumerate(DATASETS)}),
        _a=trials["arm"].map({a: i for i, a in enumerate(ARMS)}),
        _k=trials["knob"].map(lambda k: knob_rank.get(k, 99)),
        _v=pd.to_numeric(trials["value"], errors="coerce"),
    )
    trials = trials.sort_values(["_m", "_d", "_a", "seed", "_k", "_v"], kind="stable")
    front = ["model", "dataset", "arm", "seed", "knob", "value", "status", "winner"]
    metric_cols = [c for c in trials.columns if c.startswith(("val ", "test "))]
    rest = ["train (s)", "wall (h)"]
    cfg = sorted(c for c in trials.columns if c.startswith("cfg."))
    return trials[[c for c in front + metric_cols + rest + cfg if c in trials.columns]]


def _style(ws, df, row_fill=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(500))) if len(df) else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 6) + 2, 40)
        if col.startswith(("val ", "test ")) and "%" not in col:
            fmt = "0.00000"
        elif col.endswith("%"):
            fmt = "+0.0;-0.0"
        elif col in ("counterfactual_rate",):
            fmt = "0.000"
        elif col in ("train (s)",):
            fmt = "#,##0"
        elif col in ("wall (h)",):
            fmt = "0.00"
        else:
            continue
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.number_format = fmt
    if row_fill:
        for r, (_, rec) in enumerate(df.iterrows(), start=2):
            fill, bold = row_fill(rec)
            if fill or bold:
                for cell in ws[r]:
                    if fill:
                        cell.fill = fill
                    if bold:
                        cell.font = Font(bold=True)


def _trial_fill(rec):
    if rec.get("status") == "missing":
        return MISSING_FILL, False
    if rec.get("status") in ("stale", "off-grid"):
        return STALE_FILL, False
    if rec.get("winner") == "yes":
        return WIN_FILL, True
    return None, False


def main():
    if not IN_DIR.exists():
        raise SystemExit(f"no input dir: {IN_DIR}")
    trials, winners = collect()
    if trials.empty:
        raise SystemExit(f"no tuning cells in {IN_DIR}")
    trials = _order(trials)
    winners = add_arm_gap(winners)
    winners = (winners.assign(_d=winners["dataset"].map({d: i for i, d in enumerate(DATASETS)}),
                              _a=winners["arm"].map({a: i for i, a in enumerate(ARMS)}))
               .sort_values(["model", "_d", "_a", "seed"]).drop(columns=["_d", "_a"]))
    cov = coverage(winners)

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    trials.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as xw:
        trials.to_excel(xw, sheet_name="Trials", index=False)
        winners.to_excel(xw, sheet_name="Winners", index=False)
        cov.to_excel(xw, sheet_name="Coverage", index=False)
        _style(xw.sheets["Trials"], trials, _trial_fill)
        _style(xw.sheets["Winners"], winners,
               lambda r: (WIN_FILL, False) if r.get("finished") == "yes" else (None, False))
        _style(xw.sheets["Coverage"], cov)

    counts = trials["status"].value_counts().to_dict()
    print(f"[aggregate_tuning] {len(winners)} cells, trials by status: {counts}")
    print(f"[aggregate_tuning] wrote {OUT_XLSX}")
    print(f"[aggregate_tuning] wrote {OUT_CSV}")


if __name__ == "__main__":
    main()
